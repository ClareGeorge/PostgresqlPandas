import os
from sqlalchemy import create_engine, Engine
import pandas as pd
from openpyxl.workbook import Workbook

class TestSQLConcepts:

    def test_connect_to_posgresdb(self, pytestconfig ,  ):
        print( "test_sample() :" , os.environ.get('PYTEST_CURRENT_TEST'))
        print(self.postgres_db_engine)

    def test_extract_table_names_01(self):
        df_1 = pd.read_sql("select information_schema.tables.* \
                        from information_schema.tables\
                        left join information_schema.views\
                        on  information_schema.tables.table_name =  information_schema.views.table_name\
                        where information_schema.views.table_name is  null\
                        and information_schema.tables.table_schema = 'public' ", self.postgres_db_engine)
        #df_1.to_excel(os.getenv("excel_file"),"Sheet1", header=True, index=False)

        with pd.ExcelWriter(os.getenv("excel_file"),
                            mode='w',
                            engine="openpyxl",
                            ) as writer:
            df_1.to_excel(writer, sheet_name="SQL_Output", header=True, index=False)

        print(df_1)
    def test_extract_table_names_02(self):
        df_2 = pd.read_sql("select information_schema.tables.* \
                        from information_schema.tables", self.postgres_db_engine)
        df_3 = pd.read_sql("select information_schema.views.* \
                        from information_schema.views", self.postgres_db_engine)
        df_4 = pd.merge(df_2, df_3, how = "left" ,on="table_name",\
                indicator="indicator_column")
        df_4 = df_4[ (df_4['table_schema_x'] == 'public') &  (df_4["indicator_column"] == 'left_only')]
        #df_4 = pd.DataFrame(data=df_4, columns=df_2.columns)
        df_4  = df_4.drop(columns=['table_catalog_y', 'table_schema_y', 'view_definition',
                           'check_option', 'is_updatable', 'is_insertable_into_y',
                           'is_trigger_updatable', 'is_trigger_deletable',
                           'is_trigger_insertable_into', 'indicator_column'])

        df_4 = df_4.rename(columns={"table_catalog_x": "table_catalog" ,\
                                    "table_schema_x": "table_schema",\
                                    "is_insertable_into_x": "is_insertable_into"})
        with pd.ExcelWriter(os.getenv("excel_file"),
                            mode='a',
                            engine="openpyxl",
                            if_sheet_exists="replace",
                            ) as writer:
            df_4.to_excel(writer, sheet_name="DataFrame_Output", header=True, index=False)

    def test_compare_output(self):
        df_5 = pd.read_excel( open(os.getenv("excel_file"),"rb+" ),sheet_name="SQL_Output", header= 0 )
        df_6 = pd.read_excel( open(os.getenv("excel_file"), "rb+"),sheet_name= "DataFrame_Output", header= 0 )

        df_7= df_6.compare(df_5, keep_shape=True, keep_equal=True)
        df_8= df_6.compare(df_5, keep_shape=True, keep_equal=True, align_axis= 0)
        with pd.ExcelWriter(os.getenv("excel_file"),
                                mode='a',
                                engine="openpyxl",
                                if_sheet_exists="replace",
                                ) as writer:
            df_7.to_excel(writer, sheet_name="SQLvsDataFrame01", header=True, index=True)
            df_8.to_excel(writer, sheet_name="SQLvsDataFrame02", header=True, index=True)